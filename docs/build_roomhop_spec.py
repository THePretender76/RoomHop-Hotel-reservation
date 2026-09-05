from __future__ import annotations

import math
import os
import textwrap
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
OUT = DOCS / "RoomHop_Specification_Produit_Technique.docx"
ASSETS.mkdir(parents=True, exist_ok=True)

DATE_FR = "5 septembre 2026"
VERSION = "1.0"
STATUS = "Référence de conception — code validé localement, non déployé"

# standard_business_brief preset
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
TABLE_FILL = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
POSITIVE = "1F3A5F"
GOLD = "7A5A00"
RISK = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "111827"
MUTED = "5B6573"
GREEN = "276749"
AMBER_FILL = "FFF8E8"
RED_FILL = "FDECEC"
GREEN_FILL = "EAF6EF"
CYAN_FILL = "EAF6FB"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D4DAE2", size=4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_run_font(run, name="Calibri", size=11, color=BLACK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "fr-FR")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120) -> None:
    if sum(widths_dxa) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths_dxa)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_num_definitions(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    existing_abstract = [
        int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))
        if x.get(qn("w:abstractNumId"), "").isdigit()
    ]
    existing_num = [
        int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))
        if x.get(qn("w:numId"), "").isdigit()
    ]
    abstract_id = max(existing_abstract or [0]) + 1
    num_id = max(existing_num or [0]) + 1

    def make_definition(aid: int, nid: int, fmt: str, lvl_text: str, font: str | None = None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(aid))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        lvl.append(text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nid))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(aid))
        num.append(abs_ref)
        numbering.append(num)

    make_definition(abstract_id, num_id, "bullet", "•", "Symbol")
    make_definition(abstract_id + 1, num_id + 1, "decimal", "%1.")
    return num_id, num_id + 1


def assign_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167


def add_field(paragraph, code: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def fit_font(path: str, size: int):
    try:
        return ImageFont.truetype(path, size)
    except Exception:
        return ImageFont.load_default()


FONT = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"


def wrapped(draw, xy, text, font, fill, max_width, spacing=5, anchor=None):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        bbox = draw.textbbox((0, 0), candidate, font=font)
        if bbox[2] - bbox[0] <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    draw.multiline_text(xy, "\n".join(lines), font=font, fill=fill, spacing=spacing, anchor=anchor)
    return lines


def arrow(draw, start, end, fill="#506176", width=4):
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 12
    p1 = (
        end[0] - size * math.cos(angle - math.pi / 6),
        end[1] - size * math.sin(angle - math.pi / 6),
    )
    p2 = (
        end[0] - size * math.cos(angle + math.pi / 6),
        end[1] - size * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, p1, p2], fill=fill)


def box(draw, xy, title, detail="", fill="#F4F6F9", outline="#2E74B5"):
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    wrapped(draw, ((x1 + x2) / 2, y1 + 20), title, fit_font(FONT_BOLD, 23), "#0B2545", x2 - x1 - 28, anchor="ma")
    if detail:
        wrapped(draw, ((x1 + x2) / 2, y1 + 58), detail, fit_font(FONT, 17), "#4B5563", x2 - x1 - 28, anchor="ma")


def create_architecture_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 36), "Architecture logique RoomHop — état défini par le code", font=fit_font(FONT_BOLD, 36), fill="#0B2545")
    d.text((60, 84), "Flux synchrones en bleu, réplication en violet, événements en orange", font=fit_font(FONT, 20), fill="#5B6573")

    box(d, (70, 170, 320, 285), "Navigateur", "SPA React / Vite", "#EAF6FB")
    box(d, (390, 150, 670, 305), "CloudFront + WAF", "S3 privé, /api/*, /images/*", "#EAF6FB")
    box(d, (750, 150, 1040, 305), "API Gateway", "HTTP API + JWT Cognito", "#EAF6FB")
    box(d, (1120, 150, 1510, 305), "VPC Link + ALB interne", "Routage /search, /reservations, /admin", "#EAF6FB")
    arrow(d, (320, 228), (390, 228), "#2E74B5")
    arrow(d, (670, 228), (750, 228), "#2E74B5")
    arrow(d, (1040, 228), (1120, 228), "#2E74B5")

    box(d, (160, 410, 445, 555), "Search ECS", "OpenSearch prioritaire\nfallback MySQL", "#F4F6F9")
    box(d, (525, 410, 820, 555), "Reservation ECS", "Transactions, identité, événements", "#F4F6F9")
    box(d, (900, 410, 1180, 555), "RDS MySQL", "Single-AZ — source d’autorité", "#FFF8E8", "#7A5A00")
    box(d, (1250, 410, 1515, 555), "OpenSearch", "1 nœud, 1 AZ", "#F4F6F9")
    arrow(d, (1240, 305), (445, 410), "#2E74B5")
    arrow(d, (1270, 305), (675, 410), "#2E74B5")
    arrow(d, (445, 485), (900, 485), "#2E74B5")
    arrow(d, (820, 485), (900, 485), "#2E74B5")
    arrow(d, (1180, 485), (1250, 485), "#805AD5")
    d.text((1184, 448), "DMS CDC", font=fit_font(FONT_BOLD, 16), fill="#6B46C1")

    box(d, (80, 675, 370, 820), "EventBridge", "Bus + archive 365 jours", "#FFF8E8", "#D97706")
    box(d, (445, 675, 720, 820), "SQS + DLQ", "notification / analytics", "#FFF8E8", "#D97706")
    box(d, (795, 675, 1070, 820), "Lambda", "SES ou écriture S3", "#FFF8E8", "#D97706")
    box(d, (1145, 675, 1515, 820), "Données & BI", "S3 → Glue/Athena → Metabase", "#EAF6EF", "#276749")
    arrow(d, (675, 555), (225, 675), "#D97706")
    arrow(d, (370, 748), (445, 748), "#D97706")
    arrow(d, (720, 748), (795, 748), "#D97706")
    arrow(d, (1070, 748), (1145, 748), "#D97706")
    img.save(path)


def create_network_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 34), "Topologie réseau et domaines de défaillance", font=fit_font(FONT_BOLD, 36), fill="#0B2545")
    d.rounded_rectangle((80, 115, 1520, 825), radius=22, fill="#FAFBFD", outline="#805AD5", width=4)
    d.text((105, 132), "VPC 10.0.0.0/16 — aucun NAT Gateway", font=fit_font(FONT_BOLD, 24), fill="#6B46C1")
    d.rounded_rectangle((125, 205, 750, 760), radius=18, fill="#EAF6FB", outline="#2E74B5", width=3)
    d.rounded_rectangle((850, 205, 1475, 760), radius=18, fill="#EAF6FB", outline="#2E74B5", width=3)
    d.text((370, 225), "Availability Zone A", font=fit_font(FONT_BOLD, 25), fill="#0B2545")
    d.text((1095, 225), "Availability Zone B", font=fit_font(FONT_BOLD, 25), fill="#0B2545")

    box(d, (170, 300, 460, 420), "Subnet privé isolé", "ECS / endpoints", "#FFFFFF")
    box(d, (500, 300, 700, 420), "Endpoints", "11 services", "#FFFFFF")
    box(d, (895, 300, 1185, 420), "Subnet privé isolé", "ECS / endpoints", "#FFFFFF")
    box(d, (1225, 300, 1425, 420), "Endpoints", "11 services", "#FFFFFF")
    box(d, (200, 520, 470, 655), "RDS MySQL", "Single-AZ", "#FFF8E8", "#7A5A00")
    box(d, (505, 520, 705, 655), "OpenSearch", "1 nœud", "#FFF8E8", "#7A5A00")
    box(d, (925, 520, 1395, 655), "ECS services", "subnets dans 2 AZ,\ndesiredCount initial = 1", "#F4F6F9")
    arrow(d, (1170, 520), (1170, 420), "#2E74B5")
    d.text((555, 700), "La topologie couvre deux AZ, mais les composants mono-instance restent des points de défaillance.", font=fit_font(FONT_BOLD, 20), fill="#9B1C1C", anchor="ma")
    img.save(path)


def create_sequence_diagram(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "Séquence de réservation et cohérence du stock", font=fit_font(FONT_BOLD, 36), fill="#0B2545")
    lanes = [
        (180, "Client"),
        (510, "API / Reservation"),
        (900, "MySQL"),
        (1280, "EventBridge"),
    ]
    for x, label in lanes:
        d.rounded_rectangle((x - 105, 115, x + 105, 175), radius=10, fill="#E8EEF5", outline="#2E74B5", width=2)
        d.text((x, 145), label, font=fit_font(FONT_BOLD, 20), fill="#0B2545", anchor="mm")
        d.line((x, 175, x, 830), fill="#AAB4C2", width=2)

    steps = [
        (235, 180, 510, "POST + Idempotency-Key", "#2E74B5"),
        (310, 510, 900, "BEGIN ; SELECT inventaire FOR UPDATE", "#2E74B5"),
        (390, 900, 510, "lignes verrouillées + tarifs", "#2E74B5"),
        (475, 510, 900, "INSERT réservation ; UPDATE total_reserved", "#2E74B5"),
        (555, 510, 900, "COMMIT", "#276749"),
        (640, 510, 1280, "PutEvents BookingConfirmed", "#D97706"),
        (720, 1280, 510, "accepté / erreur", "#D97706"),
        (795, 510, 180, "201 ou réponse idempotente", "#2E74B5"),
    ]
    for y, x1, x2, label, color in steps:
        arrow(d, (x1, y), (x2, y), color, 4)
        d.text(((x1 + x2) / 2, y - 17), label, font=fit_font(FONT, 17), fill=color, anchor="ms")
    d.rounded_rectangle((840, 275, 970, 585), radius=12, outline="#7A5A00", width=4)
    d.multiline_text(
        (995, 430),
        "frontière\ntransactionnelle",
        font=fit_font(FONT_BOLD, 16),
        fill="#7A5A00",
        anchor="lm",
        align="left",
        spacing=2,
    )
    img.save(path)


def create_scorecard(path: Path, scores: list[tuple[str, float]]) -> None:
    img = Image.new("RGB", (1600, 760), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 35), "Maturité Well-Architected indicative — échelle 1 à 5", font=fit_font(FONT_BOLD, 36), fill="#0B2545")
    y = 130
    for label, score in scores:
        d.text((70, y + 18), label, font=fit_font(FONT_BOLD, 22), fill="#0B2545")
        x0, x1 = 520, 1450
        d.rounded_rectangle((x0, y, x1, y + 50), radius=10, fill="#E8EEF5")
        width = int((x1 - x0) * score / 5)
        color = "#2E74B5" if score >= 3.4 else ("#D97706" if score >= 2.8 else "#C53030")
        d.rounded_rectangle((x0, y, x0 + width, y + 50), radius=10, fill=color)
        d.text((1480, y + 24), f"{score:.1f}", font=fit_font(FONT_BOLD, 22), fill=color, anchor="mm")
        y += 92
    img.save(path)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Caption RoomHop" not in doc.styles:
        caption = doc.styles.add_style("Caption RoomHop", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Caption RoomHop"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = p.add_run("ROOMHOP  |  SPÉCIFICATION PRODUIT ET TECHNIQUE")
    set_run_font(left, size=8, color=MUTED, bold=True)
    right = p.add_run("\tVERSION 1.0")
    set_run_font(right, size=8, color=MUTED, bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D4DAE2")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Document de référence — non déployé  |  Page ")
    set_run_font(r, size=8, color=MUTED)
    add_field(fp, "PAGE")


def add_para(doc: Document, text: str, *, size=11, color=BLACK, bold=False, italic=False,
             after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = keep
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_label_para(doc: Document, label: str, text: str, *, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    r1 = p.add_run(label)
    set_run_font(r1, size=11, color=INK, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=BLACK)
    return p


def add_bullets(doc: Document, items: list[str], bullet_num_id: int, size=10.5, after=5):
    for item in items:
        p = doc.add_paragraph()
        assign_numbering(p, bullet_num_id)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(item)
        set_run_font(r, size=size, color=BLACK)


def add_numbered(doc: Document, items: list[str], decimal_num_id: int, size=10.3, after=5):
    for item in items:
        p = doc.add_paragraph()
        assign_numbering(p, decimal_num_id)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(item)
        set_run_font(r, size=size, color=BLACK)


def add_callout(doc: Document, title: str, text: str, *, fill=CALLOUT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent, size=8)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    t = p.add_run(title)
    set_run_font(t, size=10, color=accent, bold=True)
    body = p.add_run(f"\n{text}")
    set_run_font(body, size=10, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_source(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Preuves dans le dépôt : " + text)
    set_run_font(r, name="Consolas", size=8, color=MUTED, italic=True)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int],
              *, font_size=8.6, header_fill=TABLE_FILL, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=INK, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=BLACK, bold=(first_col_bold and idx == 0))
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code_block(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=8.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_image(doc: Document, path: Path, width: float, caption: str, alt: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", caption)
    cp = doc.add_paragraph(style="Caption RoomHop")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run(caption)


def add_page_title(doc: Document, number: str, title: str, subtitle: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(number.upper())
    set_run_font(r, size=8, color=BLUE, bold=True)
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(2)
    h.add_run(title)
    if subtitle:
        add_para(doc, subtitle, size=10, color=MUTED, italic=True, after=8)


def new_page(doc: Document, state: dict):
    if state["page"] > 0:
        doc.add_page_break()
    state["page"] += 1


def build_document() -> Path:
    arch = ASSETS / "architecture_overview.png"
    network = ASSETS / "network_topology.png"
    sequence = ASSETS / "reservation_sequence.png"
    scores = [
        ("Excellence opérationnelle", 3.2),
        ("Sécurité", 3.7),
        ("Fiabilité", 2.6),
        ("Efficacité des performances", 3.4),
        ("Optimisation des coûts", 3.0),
        ("Durabilité", 3.1),
    ]
    create_architecture_diagram(arch)
    create_network_diagram(network)
    create_sequence_diagram(sequence)
    create_scorecard(ASSETS / "well_architected_scorecard.png", scores)

    doc = Document()
    configure_styles(doc)
    configure_header_footer(doc)
    bullet_id, decimal_id = add_num_definitions(doc)
    state = {"page": 0}

    doc.core_properties.title = "RoomHop — Spécification produit et technique"
    doc.core_properties.subject = "Architecture AWS, exigences produit et évaluation Well-Architected"
    doc.core_properties.author = "RoomHop"
    doc.core_properties.keywords = "RoomHop, AWS, CDK, architecture, produit, spécification, Well-Architected"
    doc.core_properties.comments = "Document généré à partir du code local. Aucune ressource AWS déployée."

    # 1 — cover
    new_page(doc, state)
    add_para(doc, "ROOMHOP", size=12, color=BLUE, bold=True, after=58)
    add_para(doc, "SPÉCIFICATION", size=9, color=BLUE, bold=True, after=6)
    add_para(doc, "Produit et technique", size=31, color=INK, bold=True, after=8)
    add_para(
        doc,
        "Plateforme de recherche, réservation et gestion hôtelière\nArchitecture AWS définie par CDK TypeScript",
        size=15,
        color=MUTED,
        after=28,
    )
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(24)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    add_table(
        doc,
        ["Document", "Valeur"],
        [
            ["Version", VERSION],
            ["Date", DATE_FR],
            ["Statut", STATUS],
            ["Région cible", "us-east-1"],
            ["Portée", "Produit, logiciel, infrastructure, exploitation et feuille de route"],
        ],
        [2700, 6660],
        font_size=9.5,
        first_col_bold=True,
    )
    add_callout(
        doc,
        "Position de lecture",
        "Le document décrit l’état du code et des templates synthétisés. Il ne constitue ni une preuve de déploiement, ni une mesure de disponibilité, de latence ou de délivrabilité en production.",
        fill=AMBER_FILL,
        accent=GOLD,
    )

    # 2 — executive summary
    new_page(doc, state)
    add_page_title(doc, "Synthèse", "Résumé exécutif et contrôle du document")
    add_para(
        doc,
        "RoomHop est un prototype fonctionnel avancé de plateforme hôtelière. Le voyageur recherche une disponibilité réelle, réserve et annule ses séjours. Le propriétaire d’hôtel soumet une demande, attend une validation puis publie et administre son offre. Les événements métier alimentent les e-mails et l’analytique.",
    )
    add_para(
        doc,
        "L’architecture cible sépare le frontend, la recherche, les réservations, les notifications et l’analytique. Elle combine CloudFront, WAF, API Gateway, Cognito, ECS/Fargate, RDS MySQL, OpenSearch, DMS, EventBridge, SQS, Lambda, SES, S3, Glue, Athena, Metabase et un pipeline CodePipeline. Treize stacks CDK décrivent 270 ressources CloudFormation dans la dernière synthèse active.",
    )
    add_callout(
        doc,
        "Décisions intentionnelles conservées",
        "RDS MySQL reste Single-AZ. Le VPC possède deux subnets privés isolés dans deux AZ, tandis que Search, Reservation et Metabase démarrent chacun avec desiredCount = 1. Ces choix réduisent le coût d’une démonstration mais limitent la haute disponibilité.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_table(
        doc,
        ["Dimension", "Conclusion"],
        [
            ["Maturité", "Architecture de référence testée localement ; pas une plateforme exploitée."],
            ["Cohérence métier", "Transactions MySQL et verrous FOR UPDATE protègent le stock."],
            ["Résilience", "Files, DLQ, archive et fallback existent ; plusieurs composants sont mono-instance."],
            ["Sécurité", "Bon socle, mais ouverture par défaut des listeners ALB et validation des claims à renforcer."],
            ["Notification", "Chaîne codée et testée ; réception SES non vérifiable sans déploiement/configuration."],
            ["Priorité", "Fermer les listeners, renforcer validations, observabilité et atomicité des événements."],
        ],
        [2200, 7160],
        font_size=9,
        first_col_bold=True,
    )
    add_para(doc, "Public visé : sponsor produit, architecte, développeur, DevOps, sécurité, QA et futur exploitant.", size=9.5, color=MUTED, italic=True)

    # 3 — toc
    new_page(doc, state)
    add_page_title(doc, "Navigation", "Table des matières", "Pagination conçue pour une lecture dans OnlyOffice ; les numéros seront vérifiés sur le rendu final.")
    toc_rows = [
        ["1", "Vision, acteurs et périmètre", "4"],
        ["2", "Exigences fonctionnelles — voyageur", "5"],
        ["3", "Exigences fonctionnelles — partenaire et administration", "6"],
        ["4", "Parcours voyageur", "7"],
        ["5", "Parcours partenaire et notification", "8"],
        ["6", "Règles métier", "9"],
        ["7", "Exigences non fonctionnelles", "10"],
        ["8", "Architecture d’ensemble", "11"],
        ["9", "Décisions et invariants d’architecture", "12"],
        ["10", "Réseau, edge et exposition", "13"],
        ["11", "Compute, routage API et conteneurs", "14"],
        ["12", "Données relationnelles et migrations", "15"],
        ["13", "Recherche, DMS et fallback MySQL", "16"],
        ["14", "Cohérence des stocks et idempotence", "17"],
        ["15", "Événements, files et e-mails", "18"],
        ["16", "Analytique et Metabase", "19"],
        ["17", "Identité, autorisation et protection", "20"],
        ["18", "Frontend, UX et accessibilité", "21"],
        ["19", "Infrastructure as Code et CI/CD", "22"],
        ["20", "Observabilité et opérations", "23"],
        ["21", "Validation et préparation au déploiement", "24"],
        ["22", "Modèle de coût théorique", "25"],
        ["23", "Évaluation Well-Architected — méthode et synthèse", "26"],
        ["24", "Well-Architected — opérations et sécurité", "27"],
        ["25", "Well-Architected — fiabilité et performance", "28"],
        ["26", "Well-Architected — coûts et durabilité", "29"],
        ["27", "Feuille de route d’amélioration", "30"],
        ["28", "Registre des risques", "31"],
        ["29", "Traçabilité et couverture de tests", "32"],
        ["30", "Glossaire et références", "33"],
    ]
    toc_table = add_table(doc, ["§", "Section", "Page"], toc_rows, [700, 7760, 900], font_size=7.9, first_col_bold=True)
    # The TOC is intentionally compact so its explicit following page break remains
    # on this page in Word, LibreOffice and OnlyOffice-compatible layout engines.
    for row in toc_table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=45, bottom=45, start=120, end=120)

    # 4
    new_page(doc, state)
    add_page_title(doc, "1", "Vision, acteurs et périmètre")
    add_label_para(doc, "Vision. ", "Réunir découverte, disponibilité, réservation fiable, gestion partenaire et données décisionnelles dans une expérience cohérente, soutenue par une architecture AWS privée et événementielle.")
    add_table(
        doc,
        ["Acteur", "Objectif", "Capacités principales"],
        [
            ["Visiteur", "Comparer une offre", "Accueil, recherche publique, prix, disponibilités, équipements."],
            ["Voyageur", "Réserver et gérer", "Compte, réservation, historique, annulation."],
            ["Demandeur partenaire", "Obtenir un accès", "Formulaire professionnel, suivi de statut, accusé e-mail."],
            ["Partenaire approuvé", "Publier l’offre", "Création d’hôtel, chambres, tarifs, stocks, modifications API."],
            ["SuperAdmin", "Contrôler l’accès", "Lecture, approbation ou rejet des demandes."],
            ["Analyste", "Comprendre l’activité", "S3, Glue, Athena et Metabase."],
            ["Exploitant", "Livrer et maintenir", "CDK, pipeline, logs, traces, DLQ et audit."],
        ],
        [1700, 2500, 5160],
        font_size=8.7,
    )
    add_label_para(doc, "Inclus. ", "Recherche, authentification, réservation, annulation, onboarding partenaire, revue, publication d’un établissement, notifications, analytics, développement local et architecture AWS.")
    add_label_para(doc, "Hors périmètre actuel. ", "Paiement en ligne, remboursement, multi-devise, PMS/channel manager, avis clients, fidélité, application mobile, facturation PDF serveur et tableaux de bord Metabase préconfigurés.")
    add_callout(doc, "Niveau de maturité", "Prototype fonctionnel avancé / architecture de référence. Les composants sont codés et testés localement, mais aucune ressource AWS n’est déployée et aucun SLA ne peut être affirmé.")
    add_source(doc, "README.md ; infra/README.md ; infra/bin/app.ts ; hotel-ui/src/App.jsx")

    # 5
    new_page(doc, state)
    add_page_title(doc, "2", "Exigences fonctionnelles — voyageur")
    rows = [
        ["RF-G01", "Rechercher par destination, arrivée, départ et voyageurs.", "Implémenté/testé"],
        ["RF-G02", "Filtrer par prix et équipements.", "Implémenté ; étoiles partiel"],
        ["RF-G03", "Exiger tarif et stock pour chaque nuit.", "Implémenté/testé"],
        ["RF-G04", "Utiliser OpenSearch puis MySQL en cas d’erreur.", "Implémenté/testé"],
        ["RF-G05", "Créer, confirmer et ouvrir une session Cognito.", "Implémenté ; AWS non testé"],
        ["RF-G06", "Réserver sous identité JWT, jamais avec un guest_id client.", "Implémenté/testé"],
        ["RF-G07", "Calculer le montant côté serveur.", "Implémenté/testé"],
        ["RF-G08", "Éviter doublons et surréservation.", "Implémenté ; intégration réelle à ajouter"],
        ["RF-G09", "Lister et lire uniquement ses séjours.", "Implémenté/testé"],
        ["RF-G10", "Annuler et restituer le stock une seule fois.", "Implémenté/testé"],
    ]
    add_table(doc, ["ID", "Exigence", "État"], rows, [1100, 6040, 2220], font_size=8.6, first_col_bold=True)
    add_callout(
        doc,
        "Écarts fonctionnels visibles",
        "Le filtre d’étoiles est affiché mais pas appliqué. Les dates initiales de la page d’accueil sont figées en juillet 2026. L’UI réserve une chambre, ne stocke pas les demandes spéciales et le récapitulatif peut sous-estimer une période à tarifs variables.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_para(doc, "Critère d’acceptation central : deux requêtes simultanées sur la dernière unité disponible ne doivent jamais créer deux réservations confirmées.", size=10, color=INK, bold=True)
    add_source(doc, "hotel-ui/src/pages/HomePage.jsx ; SearchResultsPage.jsx ; BookingPage.jsx ; ReservationsPage.jsx ; services/search-service ; services/reservation-service")

    # 6
    new_page(doc, state)
    add_page_title(doc, "3", "Exigences fonctionnelles — partenaire et administration")
    rows = [
        ["RF-P01", "Soumettre une demande Hotel Property Owner authentifiée.", "Implémenté/testé"],
        ["RF-P02", "Ignorer toute identité Cognito fournie par le navigateur.", "Implémenté/testé"],
        ["RF-P03", "Attribuer HotelPartnerPending et exposer le statut.", "Implémenté"],
        ["RF-P04", "Envoyer l’accusé PartnerApplicationSubmitted.", "Code/test ; SES non vérifié"],
        ["RF-P05", "Limiter la décision au groupe SuperAdmin.", "Implémenté"],
        ["RF-P06", "Approuver/rejeter et synchroniser groupes + attribut.", "Implémenté/testé service"],
        ["RF-P07", "Réserver la publication aux partenaires approuvés.", "Implémenté/testé"],
        ["RF-P08", "Créer propriété, chambres, tarifs et stocks atomiquement.", "Implémenté/testé service"],
        ["RF-P09", "Limiter chaque partenaire à ses propres hôtels.", "Implémenté/testé"],
        ["RF-P10", "Administrer l’offre après création.", "API présente ; UI partielle"],
    ]
    add_table(doc, ["ID", "Exigence", "État"], rows, [1100, 6040, 2220], font_size=8.6, first_col_bold=True)
    add_bullets(doc, [
        "Champs obligatoires : identité professionnelle, société, identifiant fiscal, contact, siège, portefeuille estimé, ville principale et site HTTP/HTTPS.",
        "La création complète initialise 365 jours de tarifs et de stocks et lie l’hôtel au partenaire.",
        "L’interface de revue SuperAdmin suppose encore que l’identifiant de la demande est connu ; aucune file visuelle globale n’est finalisée.",
        "L’upload d’images et les écrans complets d’ajustement des prix/stocks restent à développer.",
    ], bullet_id, size=9.5, after=3)
    add_source(doc, "hotel-ui/src/pages/PartnerOnboardingPage.jsx ; PropertyRegistrationPage.jsx ; SuperAdminReviewPage.jsx ; services/reservation-service/src/services/adminService.js")

    # 7
    new_page(doc, state)
    add_page_title(doc, "4", "Parcours voyageur", "Recherche → authentification → réservation → gestion")
    add_numbered(doc, [
        "Le visiteur saisit destination, dates, voyageurs et prix ; les critères sont transportés dans l’URL.",
        "Le service contrôle les paramètres et recherche la disponibilité sur toute la période.",
        "OpenSearch répond en priorité. Une exception ou une absence de configuration bascule sur MySQL.",
        "Les résultats sont regroupés par hôtel ; les équipements peuvent être filtrés dans le navigateur.",
        "Le voyageur choisit un type de chambre. Une route protégée l’envoie vers la connexion si nécessaire.",
        "Le formulaire est prérempli à partir de Cognito et crée une clé d’idempotence.",
        "Le serveur rattache ou crée le client à partir du sub Cognito, verrouille chaque nuit, calcule le montant puis commit.",
        "Après le commit, BookingConfirmed est publié ; l’interface confirme et la notification est mise en file.",
        "La page Mes réservations filtre par identité Cognito. Une annulation autorisée restitue le stock dans une transaction.",
    ], decimal_id, size=9.8, after=4)
    add_callout(
        doc,
        "Règle d’annulation actuelle",
        "La fenêtre est de 72 heures après la création de la réservation. Ce n’est pas une règle « jusqu’à trois jours avant l’arrivée ». Le libellé produit doit éviter toute ambiguïté.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_table(
        doc,
        ["Réponse", "Sens"],
        [
            ["200/201", "Recherche ou réservation réussie ; replay idempotent possible."],
            ["400", "Entrée invalide ou incomplète."],
            ["401/403", "Absence d’identité ou droit insuffisant."],
            ["404", "Réservation/ressource inaccessible au compte."],
            ["409", "Stock insuffisant, période incomplète ou état incompatible."],
        ],
        [1500, 7860],
        font_size=8.8,
    )

    # 8
    new_page(doc, state)
    add_page_title(doc, "5", "Parcours partenaire et notification")
    add_numbered(doc, [
        "Un utilisateur authentifié ouvre Manage Hotel Property et complète le formulaire professionnel.",
        "Le frontend construit un payload métier sans cognitoSub ; le serveur réutilise exclusivement l’identité authentifiée.",
        "MySQL crée ou remet la demande à PENDING, puis Cognito place l’utilisateur dans HotelPartnerPending.",
        "PartnerApplicationSubmitted est envoyé au bus EventBridge.",
        "La règle partenaire alimente la file SQS de notification ; Lambda choisit le modèle et appelle SES.",
        "Le demandeur consulte une page d’attente présentant la référence et l’adresse utilisée.",
        "Un SuperAdmin approuve ou rejette. Les groupes Cognito et custom:partner_status sont mis à jour.",
        "PartnerApplicationReviewed déclenche le message de décision.",
        "Après approbation, le partenaire publie un établissement puis n’administre que les hôtels qui lui sont liés.",
    ], decimal_id, size=9.7, after=4)
    add_callout(
        doc,
        "Pourquoi un e-mail peut manquer après un futur déploiement",
        "SES exige une identité expéditrice vérifiée ; en sandbox, le destinataire doit aussi être vérifié. Il faut contrôler le paramètre SesSenderEmail, la région, les logs Lambda, la notification DLQ et l’archive EventBridge. Le code et les tests ne prouvent pas une délivrabilité réelle.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_source(doc, "infra/lib/events-stack.ts ; services/lambda/notification-handler/index.js ; services/reservation-service/src/services/adminService.js")

    # 9
    new_page(doc, state)
    add_page_title(doc, "6", "Règles métier")
    rules = [
        ["Séjour", "Intervalle semi-ouvert : arrivée incluse, départ exclu ; arrivée < départ."],
        ["Recherche", "1 à 50 voyageurs ; capacité du type de chambre ≥ voyageurs."],
        ["Disponibilité", "Une ligne de tarif et une ligne de stock doivent exister pour chaque nuit."],
        ["Stock", "Disponible = minimum de total_inventory − total_reserved sur la période."],
        ["Montant", "Somme des nightly_rate × room_count ; monnaie actuelle EUR."],
        ["Réservation", "Création en CONFIRMED ; annulation en CANCELLED, jamais suppression physique."],
        ["Identité", "sub et e-mail Cognito sont les données de confiance pour le voyageur."],
        ["Idempotence", "Une même clé ne doit produire qu’une réservation métier."],
        ["Annulation", "Propriétaire uniquement ; au plus 72 h après created_at ; stock restitué une fois."],
        ["Partenaire", "PENDING, APPROVED ou REJECTED ; décision réservée à SuperAdmin."],
        ["Propriété", "Partenaire approuvé et propriétaire du lien ; suppression logique via deleted_at."],
        ["Initialisation", "Une publication de propriété crée 365 jours de tarif et d’inventaire."],
        ["Paiement", "Pay at hotel ; aucun paiement ni remboursement en ligne."],
    ]
    add_table(doc, ["Domaine", "Règle"], rules, [1900, 7460], font_size=8.55, first_col_bold=True)
    add_callout(
        doc,
        "À durcir",
        "La validation de création de réservation vérifie surtout la présence des champs. Les types, bornes, room_count entier et positif, cohérence des dates et contraintes SQL doivent être renforcés.",
        fill=RED_FILL,
        accent=RISK,
    )
    add_source(doc, "database/migration_v2.sql ; services/reservation-service/src/middleware/validate.js ; services/reservation-service/src/services/reservationService.js")

    # 10
    new_page(doc, state)
    add_page_title(doc, "7", "Exigences non fonctionnelles")
    nfr = [
        ["Sécurité", "JWT, groupes, propriété des données, WAF, secrets et chiffrement.", "Bon socle ; chemin interne à durcir"],
        ["Cohérence", "Transactions InnoDB et FOR UPDATE.", "Fort pour stock ; tests réels à ajouter"],
        ["Résilience", "SQS/DLQ, archive, backup RDS et fallback.", "Partielle, composants mono-instance"],
        ["Performance", "CloudFront, OpenSearch, index SQL, autoscaling 1–4.", "Prometteur ; aucun SLO mesuré"],
        ["Scalabilité", "Services stateless et asynchronisme.", "Applicatif scalable ; données limitées"],
        ["Observabilité", "Logs, Container Insights, CloudTrail, daemon X-Ray.", "Alarmes/dashboard/instrumentation absents"],
        ["Maintenabilité", "Services séparés, 13 stacks, tests et pipeline.", "Bonne modularité"],
        ["Portabilité", "Environnement Docker et cible AWS.", "Risque de dérive locale/cloud"],
        ["Accessibilité", "Labels, focus, responsive, reduced motion sur onboarding.", "Audit global absent"],
        ["Confidentialité", "Chiffrement et réseau privé.", "Politique PII/RGPD à définir"],
        ["Disponibilité", "Deux AZ réseau.", "RDS, OpenSearch et tâches initiales non HA"],
        ["Coût", "Sans NAT, petites tailles, desiredCount 1.", "Endpoints privés dominants"],
    ]
    add_table(doc, ["Qualité", "Mécanisme attendu", "Appréciation"], nfr, [1600, 4800, 2960], font_size=8.1, first_col_bold=True)
    add_callout(
        doc,
        "Objectifs à formaliser",
        "Disponibilité cible, latence p50/p95, taux d’erreur, RTO, RPO, retard DMS, taux de fallback, taux de livraison e-mail et coût par réservation.",
    )

    # 11
    new_page(doc, state)
    add_page_title(doc, "8", "Architecture d’ensemble")
    add_image(
        doc,
        arch,
        6.35,
        "Figure 1 — Architecture logique définie par les stacks CDK.",
        "Diagramme de l’architecture RoomHop : CloudFront et WAF vers S3 et API Gateway, VPC Link vers ECS, RDS et OpenSearch, puis EventBridge, SQS, Lambda, SES et analytics.",
    )
    add_para(
        doc,
        "La source d’autorité transactionnelle est MySQL. OpenSearch optimise la lecture ; DMS réplique les données en full-load puis CDC. Les commandes de réservation restent synchrones, tandis que notification et analytique sont découplées par événements et files. Le frontend et les images sont privés dans S3 et servis via CloudFront.",
        size=10,
    )
    add_callout(doc, "État réel", "Le diagramme représente un end-state codé. Aucun domaine, cluster, bucket, pipeline ou base présenté ici n’existe encore dans AWS.")

    # 12
    new_page(doc, state)
    add_page_title(doc, "9", "Décisions et invariants d’architecture")
    decisions = [
        ["ADR-01", "RDS MySQL Single-AZ", "Coût et simplicité de démonstration.", "Point de panne accepté ; backup/restore impératif."],
        ["ADR-02", "Deux subnets isolés dans deux AZ", "Segmentation et préparation multi-AZ.", "Pas d’accès internet direct."],
        ["ADR-03", "desiredCount initial = 1", "Limiter le coût d’une courte exécution.", "Une tâche perdue crée une interruption jusqu’au remplacement."],
        ["ADR-04", "Aucun NAT Gateway", "Réduire coût et surface sortante.", "11 endpoints interface × 2 AZ."],
        ["ADR-05", "OpenSearch + fallback MySQL", "Recherche rapide sans abandonner l’autorité relationnelle.", "Fraîcheur DMS et résultat vide à surveiller."],
        ["ADR-06", "EventBridge + SQS", "Découpler producteur, notification et analytics.", "Cohérence commit/événement à compléter."],
        ["ADR-07", "CDK multi-stack", "Séparation des responsabilités et synthèse vérifiable.", "Ordonnancement et configuration plus complexes."],
        ["ADR-08", "Metabase conteneurisé", "BI flexible sur Athena.", "Service toujours actif, auth séparée."],
    ]
    add_table(doc, ["ID", "Décision", "Justification", "Conséquence"], decisions, [900, 2350, 2950, 3160], font_size=7.8, first_col_bold=True)
    add_callout(
        doc,
        "Invariants protégés par tests",
        "Deux subnets, zéro NAT, RDS MultiAZ=false, desiredCount=1 pour Search/Reservation, un subnet OpenSearch et Lambdas notification/analytics hors VPC.",
        fill=GREEN_FILL,
        accent=GREEN,
    )
    add_source(doc, "infra/lib/config.ts ; infra/test/architecture.test.ts")

    # 13
    new_page(doc, state)
    add_page_title(doc, "10", "Réseau, edge et exposition")
    add_image(
        doc,
        network,
        6.25,
        "Figure 2 — Répartition réseau sur deux AZ et composants mono-instance.",
        "VPC privé RoomHop avec deux subnets isolés dans deux zones, endpoints dans chaque zone, RDS et OpenSearch mono-zone, services ECS à desiredCount un.",
    )
    add_table(
        doc,
        ["Contrôle", "Configuration"],
        [
            ["Front door", "CloudFront + WAF ; buckets S3 privés via OAC. Les URLs API Gateway restent directement joignables et peuvent contourner le WAF."],
            ["API", "HTTP API, JWT Cognito, VPC Link vers ALB interne."],
            ["Sortie privée", "Endpoint S3 Gateway + 11 types d’endpoints Interface dans 2 AZ."],
            ["Security groups", "VPC Link→ALB ; ALB→ECS 3000 ; ECS/Lambda/DMS→données."],
            ["Écart critique", "Listeners ALB open par défaut : règles 0.0.0.0/0 sur 80/8080 dans le template."],
        ],
        [2100, 7260],
        font_size=8.8,
        first_col_bold=True,
    )
    add_callout(
        doc,
        "Correction recommandée avant déploiement",
        "Créer les listeners avec open: false et ne conserver que l’ingress depuis le security group VPC Link. Restreindre aussi l’accès direct aux endpoints API Gateway, limiter CORS aux origines attendues et poser CSP/HSTS sur les réponses edge.",
        fill=RED_FILL,
        accent=RISK,
    )

    # 14
    new_page(doc, state)
    add_page_title(doc, "11", "Compute, routage API et conteneurs")
    add_table(
        doc,
        ["Composant", "Taille", "Instances", "Rôle"],
        [
            ["Search ECS", "0,5 vCPU / 1 Gio", "1 initial ; autoscaling 1–4", "Recherche, OpenSearch et fallback MySQL."],
            ["Reservation ECS", "0,5 vCPU / 1 Gio", "1 initial ; autoscaling 1–4", "Réservation, annulation et admin partenaire."],
            ["Metabase ECS", "1 vCPU / 2 Gio", "1", "Exploration BI via Athena."],
            ["ALB interne", "Listeners 80 et 8080", "managé", "Routage chemins application et Metabase."],
            ["ECR", "4 repositories", "managé", "Search, Reservation, X-Ray et Metabase."],
        ],
        [2200, 1900, 2350, 2910],
        font_size=8.5,
    )
    add_label_para(doc, "Routage. ", "CloudFront envoie /api/* vers API Gateway. Les routes publiques de recherche et protégées de réservation/administration sont intégrées à l’ALB par VPC Link. Un second HTTP API et une seconde distribution exposent Metabase.")
    add_label_para(doc, "Conteneurs. ", "Les task definitions utilisent des Docker assets pour le premier déploiement ; le pipeline pousse ensuite les images versionnées dans ECR et met à jour les services ECS.")
    add_label_para(doc, "Santé. ", "Chaque service expose /health. L’ALB utilise des target groups dédiés. Les services sont stateless ; les sessions et données restent dans Cognito/MySQL.")
    add_callout(
        doc,
        "Nuance X-Ray",
        "Un sidecar daemon X-Ray est provisionné, mais aucun aws-xray-sdk n’est trouvé dans les services. Le document ne considère donc pas le traçage distribué comme opérationnel tant que l’instrumentation n’est pas ajoutée.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_source(doc, "infra/lib/compute-stack.ts ; infra/lib/api-stack.ts ; infra/lib/metabase-stack.ts ; infra/docker")

    # 15
    new_page(doc, state)
    add_page_title(doc, "12", "Données relationnelles et migrations")
    add_table(
        doc,
        ["Entité", "Responsabilité"],
        [
            ["hotel", "Métadonnées, étoiles, localisation et suppression logique."],
            ["room_type", "Capacité et équipements par catégorie."],
            ["room_type_rate", "Tarif quotidien unique par hôtel/type/date."],
            ["room_type_inventory", "Stock et réservations quotidiennes ; compteur cohérent."],
            ["guest", "Profil voyageur et liaison unique au sub Cognito."],
            ["reservation", "Séjour, statut, quantité, montant et clé d’idempotence."],
            ["hotel_images", "Références d’images et image principale."],
            ["hotel_administrators", "Demande, société, statut et identité Cognito."],
            ["hotel_admin_properties", "Relation d’appartenance partenaire-hôtel."],
        ],
        [2500, 6860],
        font_size=8.8,
        first_col_bold=True,
    )
    add_label_para(doc, "RDS. ", "MySQL 8.0.46, db.t3.medium, Single-AZ, 50 Gio extensibles à 200 Gio, chiffrement, sauvegarde sept jours, Performance Insights et secret généré.")
    add_label_para(doc, "Migration. ", "Une Lambda Node.js 24 dans les subnets isolés exécute un bootstrap idempotent et des incréments. DELETE est volontairement sans effet. La version de ressource personnalisée est 7.")
    add_label_para(doc, "Binlogs. ", "La procédure mysql.rds_set_configuration fixe la rétention à 24 heures. Cette fenêtre laisse DMS reprendre le CDC après une interruption courte ; elle doit dépasser la durée maximale de panne prévue.")
    add_callout(doc, "Limite de preuve", "Le schéma et les plans de migration sont testés, mais aucune migration n’a été exécutée contre un RDS réel.")
    add_source(doc, "infra/lib/database-stack.ts ; services/lambda/db-migration/index.js ; database/migration_v2.sql")

    # 16
    new_page(doc, state)
    add_page_title(doc, "13", "Recherche, DMS et fallback MySQL")
    add_label_para(doc, "Chemin principal. ", "Le service signe les requêtes OpenSearch avec SigV4, applique un timeout de cinq secondes, interroge des index séparés puis assemble hôtels, types, tarifs, stocks et images. La période doit être couverte nuit par nuit.")
    add_label_para(doc, "Réplication. ", "DMS dms.t3.micro, 20 Gio, Single-AZ, exécute full-load-and-cdc depuis MySQL vers OpenSearch pour hotel, room_type, room_type_rate, room_type_inventory et hotel_images.")
    add_label_para(doc, "Fallback. ", "Si l’endpoint n’est pas configuré ou si OpenSearch lève une erreur, une requête SQL équivalente est exécutée sur la source d’autorité. La réponse annonce source: opensearch ou source: mysql.")
    add_code_block(
        doc,
        "try { results = await searchOpenSearch(criteria); source = 'opensearch'; }\n"
        "catch (error) { results = await searchMySql(criteria); source = 'mysql'; }",
    )
    add_table(
        doc,
        ["Risque", "Conséquence", "Amélioration"],
        [
            ["Index incomplet mais réponse vide valide", "Aucun fallback ; faux zéro résultat.", "Vérifier fraîcheur/compteurs et fallback conditionnel."],
            ["Retard DMS > rétention binlog", "CDC irrécupérable sans rechargement.", "Alarme lag et procédure full reload."],
            ["SupportLobs=false avec hotel.description TEXT", "Description potentiellement tronquée/ignorée.", "Limited LOB ou VARCHAR borné."],
            ["Starter seulement onCreate", "Task remplacée lors d’un update peut rester STOPPED.", "Gérer onUpdate et état idempotent."],
            ["Un nœud OpenSearch", "Interruption de recherche principale.", "Accepter le fallback ou passer multi-nœud en prod."],
        ],
        [2400, 3100, 3860],
        font_size=8.0,
    )
    add_source(doc, "services/search-service/src/opensearch.js ; routes/search.js ; infra/lib/dms-stack.ts ; infra/lib/opensearch-stack.ts")

    # 17
    new_page(doc, state)
    add_page_title(doc, "14", "Cohérence des stocks et idempotence")
    add_image(
        doc,
        sequence,
        6.2,
        "Figure 3 — Transaction de réservation et publication d’événement.",
        "Diagramme de séquence : le client envoie une clé idempotente, le service verrouille l’inventaire MySQL, met à jour et commit, puis publie BookingConfirmed.",
    )
    add_para(
        doc,
        "FOR UPDATE sérialise les réservations concurrentes portant sur les mêmes lignes d’inventaire. Le service verrouille toutes les nuits dans une transaction, vérifie la couverture et le minimum disponible, écrit la réservation, incrémente total_reserved puis commit. L’annulation verrouille la réservation, contrôle son état et décrémente avec GREATEST.",
        size=9.7,
    )
    add_callout(
        doc,
        "Pourquoi les stocks restent cohérents",
        "Une deuxième transaction concurrente attend la libération des verrous, puis relit les compteurs déjà modifiés. Elle échoue si la capacité restante ne suffit plus. Rollback annule toutes les écritures en cas d’erreur.",
        fill=GREEN_FILL,
        accent=GREEN,
    )
    add_callout(
        doc,
        "Limites à traiter",
        "La contrainte d’idempotence est globale alors que la recherche est scoppée par client : une collision inter-client ou deux premières requêtes strictement simultanées peut devenir un duplicate-key/500. Ajouter gestion explicite du conflit et tests d’intégration concurrents sur MySQL réel.",
        fill=AMBER_FILL,
        accent=GOLD,
    )

    # 18
    new_page(doc, state)
    add_page_title(doc, "15", "Événements, files et e-mails")
    add_table(
        doc,
        ["Événement", "Producteur", "Consommateurs", "Effet"],
        [
            ["BookingConfirmed", "Reservation ECS", "notification + analytics", "E-mail et objet S3."],
            ["BookingCancelled", "Reservation ECS", "notification + analytics", "E-mail et objet S3."],
            ["PartnerApplicationSubmitted", "Reservation ECS", "notification", "Accusé de réception."],
            ["PartnerApplicationReviewed", "Reservation ECS", "notification", "Décision approuvée/rejetée."],
        ],
        [2700, 1900, 2400, 2360],
        font_size=8.3,
    )
    add_bullets(doc, [
        "Bus EventBridge dédié et archive 365 jours pour replay contrôlé.",
        "Une file notification et une file analytics, chacune chiffrée et associée à une DLQ ; maxReceiveCount = 3.",
        "Lambdas Node.js 24 hors VPC pour joindre SES/S3 sans NAT ; remontée des batchItemFailures message par message.",
        "Clés S3 analytiques déterministes afin qu’un retry écrase logiquement le même objet.",
    ], bullet_id, size=9.8, after=4)
    add_callout(
        doc,
        "Atomicité incomplète",
        "PutEvents intervient après COMMIT. Une panne peut laisser une réservation confirmée sans événement ; un replay idempotent peut aussi republier et doubler e-mail/analytics. Le modèle cible est une transactional outbox avec consommateur idempotent et déduplication.",
        fill=RED_FILL,
        accent=RISK,
    )
    add_callout(
        doc,
        "Exploitation SES",
        "Avant toute recette : vérifier l’expéditeur, sortir de sandbox si nécessaire, configurer DKIM/SPF/DMARC, suivre bounces/complaints et alerter sur la DLQ.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_source(doc, "infra/lib/events-stack.ts ; services/lambda/notification-handler/index.js ; analytics-handler/index.js")

    # 19
    new_page(doc, state)
    add_page_title(doc, "16", "Analytique et Metabase")
    add_numbered(doc, [
        "EventBridge duplique les événements de réservation vers la file analytique.",
        "La Lambda normalise le payload et écrit un JSON sous année/mois/jour avec une clé déterministe.",
        "Glue déclare la base et une table utilisant la projection de partitions.",
        "Athena exécute les requêtes dans un workgroup dédié et écrit ses résultats dans un bucket séparé.",
        "Metabase tourne dans ECS, utilise une base applicative dédiée initialisée par Lambda et accède à Athena via son rôle IAM.",
        "Une distribution CloudFront et une HTTP API distinctes exposent l’interface Metabase via un listener ALB 8080.",
    ], decimal_id, size=9.8, after=5)
    add_table(
        doc,
        ["Atout", "Limite actuelle"],
        [
            ["Découplage du transactionnel", "Seulement BookingConfirmed/Cancelled sont analysés."],
            ["Partitionnement sans crawler obligatoire", "Format JSON moins compact que Parquet."],
            ["Permissions via rôle de tâche", "Connexion Athena et dashboards à configurer manuellement."],
            ["Base Metabase séparée", "Authentification Metabase indépendante de Cognito."],
            ["Service privé derrière edge", "desiredCount 1 ; l’HTTP API directe peut contourner le WAF."],
        ],
        [4300, 5060],
        font_size=8.8,
    )
    add_callout(doc, "Évolution recommandée", "Passer à Parquet, définir un schéma versionné et des métriques certifiées. Réduire ou pseudonymiser les PII dans les événements, borner la rétention après Glacier et remplacer les permissions Athena/Glue génériques par le moindre privilège.")
    add_source(doc, "infra/lib/analytics-stack.ts ; infra/lib/metabase-stack.ts ; services/lambda/analytics-handler/index.js")

    # 20
    new_page(doc, state)
    add_page_title(doc, "17", "Identité, autorisation et protection")
    add_table(
        doc,
        ["Couche", "Contrôle"],
        [
            ["Identité utilisateur", "Cognito User Pool, client SPA, confirmation e-mail, TOTP optionnel."],
            ["Rôles métier", "Guest, HotelPartnerPending, HotelPartner, SuperAdmin."],
            ["API edge", "Authorizer JWT sur routes protégées ; recherche publique ; CORS actuellement permissif (*)."],
            ["Autorisation service", "Contrôle groupe, statut partenaire et propriété hotel/réservation."],
            ["Secrets", "Secrets Manager pour MySQL et Metabase ; injection par rôle d’exécution."],
            ["Données", "Chiffrement RDS/OpenSearch/S3/SQS/ECR ; TLS et subnets privés."],
            ["Protection edge", "WAF : rate-limit, Common Rule Set et SQLi Rule Set."],
            ["Audit", "CloudTrail single-region + S3/Logs ; IAM Access Analyzer compte."],
        ],
        [2350, 7010],
        font_size=8.8,
        first_col_bold=True,
    )
    add_callout(
        doc,
        "Risque de confiance réseau",
        "Le middleware Reservation accepte les claims transmis par le chemin interne sans vérifier lui-même une signature JWT. Tant que les listeners restent ouverts au VPC, un appel interne peut contourner l’authorizer. Fermer l’ALB et valider cryptographiquement le jeton ou un mécanisme d’authentification de service.",
        fill=RED_FILL,
        accent=RISK,
    )
    add_bullets(doc, [
        "Activer rotation des secrets et procédures d’urgence.",
        "Ajouter GuardDuty, Security Hub, Inspector/ECR scanning et cdk-nag.",
        "Journaliser les décisions admin et protéger les données personnelles par politique RGPD.",
    ], bullet_id, size=9.5, after=3)
    add_source(doc, "infra/lib/auth-stack.ts ; frontend-stack.ts ; observability-stack.ts ; services/reservation-service/src/middleware/auth.js")

    # 21
    new_page(doc, state)
    add_page_title(doc, "18", "Frontend, UX et accessibilité")
    screenshot = ROOT / "screenshot-test" / "Screenshot 2026-06-30 191442.png"
    if screenshot.exists():
        add_image(
            doc,
            screenshot,
            6.2,
            "Figure 4 — Écran de recherche servant de référence visuelle RoomHop.",
            "Capture de l’écran RoomHop Search Hotels avec filtres, cartes d’hôtels, types de chambre, disponibilité et tarif.",
        )
    add_para(
        doc,
        "La SPA React/Vite applique une navigation publique et protégée. AuthContext centralise Cognito, useAuth évite les dépendances circulaires, le client HTTP ajoute le token et PrivateRoute oriente les statuts partenaire. CloudFront sert les routes SPA, l’API et les images sous une origine cohérente.",
        size=9.5,
    )
    add_callout(
        doc,
        "Onboarding propriétaire harmonisé",
        "Le formulaire professionnel reprend la palette bleu/gris, les cartes, sections numérotées, espacements, états d’erreur/chargement, focus clavier, responsive et préférence reduced motion déjà présents dans recherche et authentification.",
        fill=GREEN_FILL,
        accent=GREEN,
    )
    add_bullets(doc, [
        "Restent à traiter : dates dynamiques, filtre étoiles, cohérence du résumé tarifaire et écrans partenaire complets.",
        "Un audit WCAG global, des tests clavier/lecteur d’écran et des tests visuels navigateur sont encore requis.",
    ], bullet_id, size=9.2, after=3)
    add_source(doc, "hotel-ui/src/App.jsx ; auth ; pages ; PartnerOnboardingPage.module.css ; __tests__")

    # 22
    new_page(doc, state)
    add_page_title(doc, "19", "Infrastructure as Code et CI/CD")
    add_table(
        doc,
        ["Étape", "Service", "Résultat"],
        [
            ["Source", "GitHub via CodeConnections", "Déclenchement sur branche paramétrée ; autorisation manuelle initiale."],
            ["Validate", "CodeBuild", "Tests, lint, build et synthèse CDK."],
            ["ContainerBuild", "CodeBuild Docker", "Images Search, Reservation, X-Ray et Metabase poussées dans ECR."],
            ["FrontendBuild", "CodeBuild", "Build Vite configuré avec outputs Cognito/API."],
            ["Deploy", "ECS + S3 actions", "Trois services mis à jour ; frontend copié dans S3."],
            ["Invalidate", "CodeBuild", "Invalidation CloudFront après livraison frontend."],
        ],
        [1600, 2500, 5260],
        font_size=8.7,
    )
    add_label_para(doc, "CDK. ", "Treize stacks TypeScript séparent Network, Database, OpenSearch, Auth, Events, Compute, Api, Frontend, Analytics, Metabase, DMS, Observability et Pipeline.")
    add_label_para(doc, "Premier déploiement. ", "Les ECS tasks utilisent des Docker assets CDK ; elles ne dépendent pas d’images préexistantes dans ECR.")
    add_callout(
        doc,
        "Limite du pipeline",
        "La phase Validate synthétise l’infrastructure mais le pipeline ne lance pas cdk deploy. Il livre uniquement images ECS, assets S3 et invalidation. Les évolutions de stacks nécessitent donc un workflow d’infrastructure séparé, idéalement avec diff, approbation et déploiement contrôlé.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_bullets(doc, [
        "Ajouter cdk diff, cdk-nag, scans dépendances/images et tests d’intégration.",
        "Prévoir comptes dev/stage/prod, approbation avant production et stratégie canary/blue-green.",
        "Versionner schémas et événements avec critères de compatibilité.",
    ], bullet_id, size=9.4, after=3)
    add_source(doc, "infra/lib/pipeline-stack.ts ; infra/bin/app.ts ; infra/package.json")

    # 23
    new_page(doc, state)
    add_page_title(doc, "20", "Observabilité et opérations")
    add_table(
        doc,
        ["Signal", "Présent", "Manquant / amélioration"],
        [
            ["Logs applicatifs", "Groupes CloudWatch explicites, logs structurés.", "Corrélation uniforme et politique PII."],
            ["Conteneurs", "Container Insights.", "Alertes CPU/mémoire/restarts/targets."],
            ["Traçage", "Daemon X-Ray sidecar.", "Instrumentation SDK et propagation des segments."],
            ["Audit", "CloudTrail management events vers S3 + Logs.", "Trail multi-région, alertes sensibles et expiration des versions S3 non courantes."],
            ["Accès", "IAM Access Analyzer.", "Traitement formalisé des findings."],
            ["Asynchrone", "DLQ et archive EventBridge.", "Alarmes âge/profondeur/DLQ et runbook replay."],
            ["Données", "Backup RDS 7 jours, Performance Insights.", "Test restore, lag DMS, espace, connexions."],
            ["Expérience", "Aucun SLO ni canary.", "Synthetics, p95, taux d’erreur et parcours critique."],
        ],
        [1700, 3300, 4360],
        font_size=8.2,
    )
    add_callout(
        doc,
        "À ne pas surévaluer",
        "La stack Observability ne crée ni dashboard ni alarmes métier. Le document considère donc l’observabilité comme un socle de télémétrie, pas comme une capacité d’exploitation aboutie.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_bullets(doc, [
        "Définir SLI/SLO pour recherche, réservation et e-mail.",
        "Créer runbooks : stock conflictuel, DMS lag, OpenSearch indisponible, SES bounce, rollback ECS et restauration RDS.",
        "Tester périodiquement alertes, replay et restauration.",
    ], bullet_id, size=9.7, after=4)
    add_source(doc, "infra/lib/observability-stack.ts ; logs dans compute/events/database/metabase ; aucun AWS::CloudWatch::Alarm dans la synthèse")

    # 24
    new_page(doc, state)
    add_page_title(doc, "21", "Validation et préparation au déploiement")
    add_table(
        doc,
        ["Validation", "Résultat observé"],
        [
            ["Backend et services", "65/65 tests réussis."],
            ["Frontend", "10/10 tests, lint et build réussis."],
            ["Infrastructure", "5/5 tests d’architecture, TypeScript et synthèse de 13 stacks réussis."],
            ["Conteneurs", "4/4 images construites localement."],
            ["AWS", "Aucun déploiement ; Cognito, SES, DMS, OpenSearch, réseau et pipeline non testés intégrés."],
        ],
        [2800, 6560],
        font_size=9.2,
        first_col_bold=True,
    )
    add_para(doc, "Checklist avant une première recette éphémère :", size=10.5, color=INK, bold=True)
    add_bullets(doc, [
        "Définir compte/région, bootstrap CDK et tags ; examiner cdk diff et coûts.",
        "Fermer les listeners ALB, renforcer JWT/validation serveur et corriger dates/filtre étoiles.",
        "Vérifier SesSenderEmail, destinataires sandbox, DKIM et configuration des événements.",
        "Autoriser CodeConnections GitHub et décider qui déploie réellement les stacks.",
        "Charger les images, initialiser Metabase, connecter Athena et restreindre son accès.",
        "Exécuter tests E2E, concurrence MySQL réelle, panne OpenSearch, DLQ/replay et restauration.",
        "Créer Budget/Anomaly Detection et une procédure de destruction ; contrôler les ressources RETAIN.",
    ], bullet_id, size=9.4, after=3)
    add_callout(
        doc,
        "Commande de synthèse",
        "cd infra ; npm ci ; npm run build ; npm test ; npm run synth. La synthèse produit des templates ; elle ne crée aucune ressource AWS.",
        fill=CALLOUT,
        accent=BLUE,
    )
    add_source(doc, "résultats de validation du projet ; infra/README.md ; package.json des composants")

    # 25
    new_page(doc, state)
    add_page_title(doc, "22", "Modèle de coût théorique", "Ordre de grandeur à la demande, us-east-1, daté du 5 septembre 2026")
    rates = [
        ("11 endpoints Interface dans 2 AZ", 0.2200),
        ("Fargate : Search + Reservation + Metabase", 0.0987),
        ("RDS MySQL db.t3.medium Single-AZ", 0.0680),
        ("OpenSearch t3.small.search, 1 nœud", 0.0360),
        ("DMS dms.t3.micro", 0.0360),
        ("ALB interne, composante horaire", 0.0225),
        ("Stockage provisionné agrégé, ordre de grandeur", 0.0108),
    ]
    hourly = sum(v for _, v in rates)
    assert round(hourly, 4) == 0.4920
    four_hours = hourly * 4
    five_hours = hourly * 5
    cost_rows = [[label, f"{value:.4f} USD/h", f"{value / hourly * 100:.1f} %"] for label, value in rates]
    add_table(doc, ["Poste fixe", "Taux estimé", "Part"], cost_rows, [5600, 1900, 1860], font_size=8.6)
    add_table(
        doc,
        ["Durée", "Socle fixe", "Avec faible trafic/logs"],
        [
            ["4 heures", f"{four_hours:.2f} USD", "≈ 2,03 USD"],
            ["5 heures", f"{five_hours:.2f} USD", "≈ 2,52 USD"],
            ["1 exécution pipeline", "hors socle", "≈ +0,39 USD"],
        ],
        [2500, 3000, 3860],
        font_size=9.2,
    )
    add_callout(
        doc,
        "Lecture correcte",
        "Estimation brute hors taxes, transferts importants, requêtes, logs volumineux et variations tarifaires. Les endpoints privés dominent le socle. OpenSearch en RemovalPolicy RETAIN peut continuer à coûter environ 0,94 USD/jour s’il est oublié après destruction.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Référence : ")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    add_hyperlink(p, "AWS Pricing Calculator", "https://calculator.aws/")

    # 26
    new_page(doc, state)
    add_page_title(doc, "23", "Évaluation Well-Architected — méthode et synthèse")
    add_para(
        doc,
        "Cette évaluation interne confronte le code, les templates synthétisés et les tests aux six piliers AWS. Elle n’est pas une revue officielle réalisée dans AWS Well-Architected Tool. En l’absence de déploiement, elle évalue l’intention et les garde-fous, pas les données d’exploitation.",
        size=10,
    )
    add_image(
        doc,
        ASSETS / "well_architected_scorecard.png",
        6.15,
        "Figure 5 — Score de maturité indicatif par pilier.",
        "Barres de score Well-Architected sur cinq : opérations 3,2 ; sécurité 3,7 ; fiabilité 2,6 ; performance 3,4 ; coûts 3,0 ; durabilité 3,1.",
    )
    overall = sum(v for _, v in scores) / len(scores)
    add_table(
        doc,
        ["Échelle", "Interprétation"],
        [
            ["1", "Risque élevé ou absence de mécanisme."],
            ["2", "Socle partiel ; lacunes structurantes."],
            ["3", "Conception cohérente ; exploitation à compléter."],
            ["4", "Pratique solide, mesurée et automatisée."],
            ["5", "Optimisation continue démontrée."],
        ],
        [1200, 8160],
        font_size=8.8,
    )
    add_callout(doc, "Résultat global", f"Moyenne indicative {overall:.1f}/5. La sécurité et la performance ont un bon socle ; la fiabilité est volontairement limitée par les choix mono-instance.")

    # 27
    new_page(doc, state)
    add_page_title(doc, "24", "Well-Architected — excellence opérationnelle et sécurité")
    add_table(
        doc,
        ["Pilier", "Forces", "Écarts prioritaires", "Score"],
        [
            [
                "Excellence opérationnelle",
                "CDK, 13 stacks, tests, pipeline V2, logs explicites, CloudTrail, DLQ/archive.",
                "Pas de déploiement infra dans pipeline, dashboard/alarme/SLO/runbooks ; X-Ray non instrumenté ; documentation locale divergente.",
                "3,2/5",
            ],
            [
                "Sécurité",
                "Cognito, groupes, JWT edge, propriété métier, WAF, chiffrement, secrets, réseau isolé, Access Analyzer.",
                "Listeners ALB ouverts au VPC, confiance claims interne, rotation/GuardDuty/Security Hub/cdk-nag/PII à ajouter.",
                "3,7/5",
            ],
        ],
        [1750, 2750, 3900, 960],
        font_size=8.4,
    )
    add_para(doc, "Appréciation détaillée — opérations", size=11, color=INK, bold=True)
    add_bullets(doc, [
        "Prepare : infrastructure et tests automatisés, mais readiness review, runbooks et RTO/RPO non formalisés.",
        "Operate : télémétrie créée, mais aucune alarme ni procédure de triage codée.",
        "Evolve : pipeline favorise les changements applicatifs ; le chemin d’évolution des stacks reste manuel.",
    ], bullet_id, size=9.5, after=3)
    add_para(doc, "Appréciation détaillée — sécurité", size=11, color=INK, bold=True)
    add_bullets(doc, [
        "Identité forte au périmètre utilisateur et contrôles d’appartenance côté service.",
        "Défense en profondeur incomplète entre API Gateway et ALB tant que open: false et validation cryptographique ne sont pas appliqués.",
        "Protection des données convenable par défaut ; gouvernance, rotation et réponse à incident à industrialiser.",
    ], bullet_id, size=9.5, after=3)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_hyperlink(p, "AWS — Operational Excellence Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/operational-excellence-pillar/welcome.html")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    add_hyperlink(p2, "AWS — Security Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/security-pillar/welcome.html")

    # 28
    new_page(doc, state)
    add_page_title(doc, "25", "Well-Architected — fiabilité et efficacité des performances")
    add_table(
        doc,
        ["Pilier", "Forces", "Écarts prioritaires", "Score"],
        [
            [
                "Fiabilité",
                "Transactions, idempotence, fallback, autoscaling, SQS/DLQ, archive, sauvegarde et IaC.",
                "RDS Single-AZ, desiredCount 1, OpenSearch 1 nœud, DMS Single-AZ, pas d’outbox ni tests de reprise/charge.",
                "2,6/5",
            ],
            [
                "Performance",
                "CloudFront, OpenSearch, indexes SQL, services stateless, autoscaling 1–4, traitement asynchrone.",
                "Aucune mesure p95/charge, assemblage multi-index, résultats vides non contrôlés, tailles non benchmarkées.",
                "3,4/5",
            ],
        ],
        [1500, 2850, 4050, 960],
        font_size=8.4,
    )
    add_callout(
        doc,
        "Fiabilité sous contraintes",
        "Le maintien volontaire de Single-AZ et desiredCount 1 est compatible avec une démonstration de quelques heures, pas avec un objectif de haute disponibilité. Tant que ces invariants restent imposés, les compensations sont : backups, restore drill, fallback testé, redeploy rapide, alerte et communication d’indisponibilité.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    add_bullets(doc, [
        "Définir RTO/RPO et tester perte RDS, panne OpenSearch, arrêt task ECS, lag DMS et DLQ.",
        "Ajouter transactional outbox et déduplication pour la chaîne après commit.",
        "Mesurer recherche/réservation p50/p95, taux 4xx/5xx, connexions DB, CPU/mémoire, cache CloudFront et lag.",
        "Benchmark avant redimensionnement ; tester montée 1→4 tasks et limites MySQL.",
    ], bullet_id, size=9.5, after=4)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_hyperlink(p, "AWS — Reliability Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/reliability-pillar/welcome.html")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    add_hyperlink(p2, "AWS — Performance Efficiency Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/performance-efficiency-pillar/welcome.html")

    # 29
    new_page(doc, state)
    add_page_title(doc, "26", "Well-Architected — optimisation des coûts et durabilité")
    add_table(
        doc,
        ["Pilier", "Forces", "Écarts prioritaires", "Score"],
        [
            [
                "Optimisation des coûts",
                "Pas de NAT, petites tailles, desiredCount 1, autoscaling, services managés et estimation éphémère.",
                "22 endpoints-AZ, services always-on, aucun budget/tagging/plan d’arrêt, RETAIN OpenSearch et droitsizing non mesuré.",
                "3,0/5",
            ],
            [
                "Durabilité",
                "Services managés, CDN, asynchronisme, petites capacités et absence de NAT.",
                "Aucun KPI d’impact, cycle de vie data, droitsizing réel, Graviton ni suppression automatique des environnements.",
                "3,1/5",
            ],
        ],
        [1650, 2800, 3950, 960],
        font_size=8.4,
    )
    add_bullets(doc, [
        "Créer tags de coûts obligatoires, AWS Budget, Anomaly Detection et alerte avant une démonstration.",
        "Évaluer si les 11 endpoints doivent tous exister dans deux AZ pour un environnement court ; documenter le compromis coût/résilience.",
        "Automatiser arrêt/destruction et inventaire des ressources RETAIN ; vérifier après cdk destroy.",
        "Mesurer coût, vCPU-seconde, Gio-stockés et données transférées par recherche/réservation.",
        "Compacter l’analytics en Parquet, appliquer lifecycle et supprimer les données de test.",
        "Comparer Fargate ARM64/Graviton et versions logicielles plus efficaces après benchmark.",
    ], bullet_id, size=9.7, after=4)
    add_callout(
        doc,
        "Compromis central",
        "Les endpoints privés suppriment le NAT et renforcent le contrôle réseau, mais constituent environ 45 % du socle estimé. La bonne décision dépend de la durée, du niveau d’isolation et du besoin multi-AZ.",
        fill=AMBER_FILL,
        accent=GOLD,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_hyperlink(p, "AWS — Cost Optimization Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/cost-optimization-pillar/welcome.html")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    add_hyperlink(p2, "AWS — Sustainability Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/sustainability-pillar/sustainability-pillar.html")

    # 30
    new_page(doc, state)
    add_page_title(doc, "27", "Feuille de route d’amélioration")
    roadmap = [
        ["P0 — avant déploiement", "Fermer listeners ALB ; validation JWT/service ; dates dynamiques ; filtre étoiles ; validation serveur stricte ; contraintes DB.", "Sécurité et exactitude produit."],
        ["P0 — fiabilité métier", "Transactional outbox, déduplication consommateurs, gestion duplicate idempotency, tests MySQL concurrents.", "Éviter événement perdu/doublé et erreurs 500."],
        ["P0 — exploitation", "Alarmes, dashboard, SLO, logs corrélés, DLQ/replay, DMS lag, SES bounce/complaint et runbooks.", "Détecter et restaurer rapidement."],
        ["P0 — gouvernance", "Tags, Budget, Anomaly Detection, procédure destroy/RETAIN, secrets et PII.", "Limiter coût et risque."],
        ["P1 — expérience", "Montant exact avant confirmation, room_count, demandes spéciales, gestion partenaire complète, file SuperAdmin, upload images.", "Finir les parcours."],
        ["P1 — recherche", "Fraîcheur d’index, fallback sur incohérence, Support LOB DMS, starter onUpdate, alarmes de lag.", "Réduire faux zéros."],
        ["P1 — livraison", "Pipeline infra séparé avec diff/approbation, scans, comptes dev/stage/prod et E2E.", "Changements contrôlés."],
        ["P2 — production", "Si contraintes levées : RDS Multi-AZ, desiredCount ≥2, OpenSearch multi-nœud/zone, stratégie DR.", "Haute disponibilité."],
        ["P2 — data", "Parquet, métriques certifiées, lifecycle, privacy retention et dashboards préconfigurés.", "BI économique et gouvernée."],
    ]
    add_table(doc, ["Horizon", "Actions", "Valeur"], roadmap, [2050, 5340, 1970], font_size=8.0, first_col_bold=True)
    add_callout(
        doc,
        "Ordre recommandé",
        "Corriger d’abord les frontières de confiance et l’atomicité. Ajouter ensuite mesure et tests intégrés. Les changements de haute disponibilité restent conditionnels : ils ne modifient pas les deux invariants explicitement conservés pour la version actuelle.",
        fill=GREEN_FILL,
        accent=GREEN,
    )
    add_para(doc, "Definition of Done P0 : synthèse propre, tests verts, aucune route interne contournable, réservation/événement récupérables, alertes testées et destruction contrôlée.", size=9.8, color=INK, bold=True)

    # 31
    new_page(doc, state)
    add_page_title(doc, "28", "Registre des risques")
    risks = [
        ["R-01", "Listener ALB ouvert par défaut", "Élevé", "Élevé", "open:false + ingress SG VPC Link"],
        ["R-02", "Claims internes non vérifiés cryptographiquement", "Moyen", "Élevé", "JWT/service auth + réseau fermé"],
        ["R-03", "Commit sans événement ou doublon au replay", "Moyen", "Élevé", "Outbox + déduplication"],
        ["R-04", "OpenSearch vide mais index incomplet", "Moyen", "Moyen", "Fraîcheur + fallback contrôlé"],
        ["R-05", "RDS Single-AZ indisponible", "Moyen", "Élevé", "Backup, restore drill, RTO/RPO ; Multi-AZ en prod"],
        ["R-06", "Une seule task applicative", "Moyen", "Moyen", "Auto-remplacement ; desired≥2 en prod"],
        ["R-07", "OpenSearch/DMS mono-instance", "Moyen", "Moyen", "Fallback, alarmes, procédure reload"],
        ["R-08", "SES sandbox/identité non prête", "Élevé", "Moyen", "Recette SES et suivi DLQ"],
        ["R-09", "Pipeline ne déploie pas l’IaC", "Élevé", "Moyen", "Workflow infra dédié"],
        ["R-10", "Divergence local/AWS et README obsolète", "Élevé", "Faible", "Docs générées, tests contractuels"],
        ["R-11", "Validation réservation permissive", "Moyen", "Élevé", "Schémas stricts + contraintes SQL"],
        ["R-12", "PII/RGPD non gouverné", "Moyen", "Élevé", "Classification, rétention, droits d’accès"],
    ]
    add_table(doc, ["ID", "Risque", "Prob.", "Impact", "Traitement"], risks, [800, 3230, 900, 900, 3530], font_size=7.8, first_col_bold=True)
    add_callout(
        doc,
        "Risque résiduel accepté",
        "R-05 et R-06 ne sont pas supprimés dans la version actuelle : Single-AZ et desiredCount 1 sont des contraintes volontaires. Leur acceptation doit être limitée à un environnement de démonstration sans engagement de disponibilité.",
        fill=AMBER_FILL,
        accent=GOLD,
    )

    # 32
    new_page(doc, state)
    add_page_title(doc, "29", "Traçabilité et couverture de tests")
    coverage = [
        ["Recherche", "Paramètres, prix, disponibilité, fallback", "5 tests service + propriétés API", "Unitaire/contractuel"],
        ["Réservation", "Montant, stock, idempotence, propriété", "8 tests service + propriétés API", "Mocks ; MySQL réel absent"],
        ["Partenaire", "Identité, validation, statut, événements", "Tests admin + 2 tests UI", "Service/UI"],
        ["Notification", "4 modèles, parsing, SES, batch failures", "5 Lambda + 6 local", "Client SES simulé"],
        ["Analytics", "Normalisation et clé S3", "2 tests", "Client S3 simulé"],
        ["Migration", "Plan bootstrap/incréments/DELETE", "4 tests", "Pas de RDS réel"],
        ["Metabase init", "Création DB/utilisateur", "1 test", "Connexion simulée"],
        ["Frontend", "Formulaire et idempotency hook", "10 tests + lint/build", "Pas d’E2E navigateur final"],
        ["CDK", "Invariants réseau/RDS/ECS/Lambda", "5 tests + tsc + synth", "13 stacks, 270 ressources"],
        ["Docker", "Build images", "4/4", "Pas de scan/registry"],
    ]
    add_table(doc, ["Domaine", "Comportements couverts", "Preuve", "Limite"], coverage, [1500, 3300, 2300, 2260], font_size=7.9)
    add_callout(
        doc,
        "Bilan",
        "65 tests backend/services et 10 tests frontend réussissent. Cette couverture réduit le risque de régression locale, mais ne valide ni IAM, ni réseau, ni intégrations managées, ni performance, ni récupération.",
        fill=GREEN_FILL,
        accent=GREEN,
    )
    add_bullets(doc, [
        "Ajouter tests de concurrence contre MySQL 8 réel.",
        "Ajouter E2E Cognito→API→ECS→RDS et EventBridge→SQS→Lambda→SES.",
        "Ajouter tests de panne, charge, accessibilité et sécurité des frontières.",
    ], bullet_id, size=9.5, after=3)
    add_source(doc, "hotel-api/__tests__ ; hotel-ui/__tests__ ; services/*/test ; services/lambda/*/test ; infra/test/architecture.test.ts")

    # 33
    new_page(doc, state)
    add_page_title(doc, "30", "Glossaire et références")
    glossary = [
        ["AZ", "Availability Zone, domaine de défaillance isolé dans une région."],
        ["CDC", "Change Data Capture, réplication des changements depuis les binlogs."],
        ["CDK", "Cloud Development Kit, infrastructure décrite en TypeScript."],
        ["DLQ", "Dead-letter queue, file recevant les messages après échecs répétés."],
        ["Idempotence", "Même demande répétée : un seul effet métier."],
        ["OAC", "Origin Access Control, accès CloudFront à un bucket S3 privé."],
        ["RPO / RTO", "Perte de données tolérée / délai de restauration toléré."],
        ["SLO", "Objectif mesurable de niveau de service."],
        ["SigV4", "Signature AWS des requêtes API."],
        ["Transactional outbox", "Événement écrit avec la transaction puis publié de façon fiable."],
    ]
    add_table(doc, ["Terme", "Définition"], glossary, [2000, 7360], font_size=8.4, first_col_bold=True)
    add_para(doc, "Références AWS officielles", size=11, color=INK, bold=True, after=3)
    links = [
        ("AWS Well-Architected — vue d’ensemble", "https://aws.amazon.com/architecture/well-architected/"),
        ("Operational Excellence Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/operational-excellence-pillar/welcome.html"),
        ("Security Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/security-pillar/welcome.html"),
        ("Reliability Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/reliability-pillar/welcome.html"),
        ("Performance Efficiency Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/performance-efficiency-pillar/welcome.html"),
        ("Cost Optimization Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/cost-optimization-pillar/welcome.html"),
        ("Sustainability Pillar", "https://docs.aws.amazon.com/wellarchitected/latest/sustainability-pillar/sustainability-pillar.html"),
        ("AWS Pricing Calculator", "https://calculator.aws/"),
    ]
    for label, url in links:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_hyperlink(p, label, url)
    add_para(doc, "Références projet principales", size=11, color=INK, bold=True, after=3)
    add_para(
        doc,
        "README.md ; infra/README.md ; infra/bin/app.ts ; infra/lib/*.ts ; infra/test/architecture.test.ts ; database/migration_v2.sql ; hotel-ui/src ; services/search-service ; services/reservation-service ; services/lambda.",
        size=8.6,
        color=MUTED,
    )
    add_callout(doc, "Fin du document", "La prochaine révision doit intégrer les résultats d’un déploiement de recette, les mesures de charge, les preuves SES/DMS et les décisions prises sur les améliorations P0.")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    output = build_document()
    print(output)
